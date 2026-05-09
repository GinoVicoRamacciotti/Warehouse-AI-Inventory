import pandas as pd
from openpyxl import load_workbook
import config

def cargar_sap():
    """
    Lee el archivo Excel de SAP y devuelve un DataFrame y un diccionario
    con el mapa de HU -> Ubicación esperado.
    """
    try:
        df_sap = pd.read_excel(config.ARCHIVO_SAP, sheet_name=config.NOMBRE_HOJA, dtype=str)
        df_sap.columns = [c.strip() for c in df_sap.columns]
        df_sap[config.COL_UBICACION] = df_sap[config.COL_UBICACION].astype(str).str.strip().str.upper()

        sap_hu_map = {}
        for _, row in df_sap.iterrows():
            hu_val = str(row[config.COL_HU]).strip().replace(".0", "")
            ubi_val = str(row[config.COL_UBICACION]).strip().upper()
            if hu_val and hu_val not in ["nan", "None", "", "<< vacías >>"]:
                sap_hu_map[hu_val] = ubi_val
        
        return df_sap, sap_hu_map
    except Exception as e:
        print(f"[ERROR] Error SAP: {e}")
        return None, {}

def determinar_estado(ubi_fisica, hu_fisica, df_sap):
    """
    Determina el estado final a guardar en el SAP basándose en la Ubicación y HU leídas.
    """
    if ubi_fisica == "NO_LEIDO": 
        return "ERROR VISUAL: ILEGIBLE"
        
    fila = df_sap[df_sap[config.COL_UBICACION] == ubi_fisica]
    if fila.empty: 
        return f"ERROR: No existe en SAP"
        
    sap_hu = str(fila.iloc[0][config.COL_HU]).strip().replace(".0", "")
    sap_vacio = (sap_hu in ["nan", "", "None", "NaN", "<< vacías >>"])
    
    if hu_fisica == "VACIO": 
        return "OK" if sap_vacio else "A - Físico Vacío / SAP Ocupado"
        
    if hu_fisica == "NO_LEIDO": 
        return "E - HU Ilegible"
        
    if not sap_vacio:
        if sap_hu.endswith(str(hu_fisica)[-6:]): 
            return "OK"
        else: 
            return "C - Diferencia HU"
            
    return "B - Físico Ocupado / SAP Vacío"

def guardar_reportes_temporales(registros_globales, reporte_relaciones):
    """
    Guarda los dataframes de registros y relaciones en Excel.
    """
    if registros_globales:
        pd.DataFrame(registros_globales).to_excel(config.ARCHIVO_REPORTE_DETALLADO, index=False)
    if reporte_relaciones:
        pd.DataFrame(reporte_relaciones).to_excel(config.ARCHIVO_RELACIONES, index=False)

def actualizar_excel_sap(resultados_finales):
    """
    Actualiza directamente el archivo Excel de SAP con los estados de resultados_finales.
    resultados_finales = { "UBI1": {"status": "OK"}, ... }
    """
    try:
        wb = load_workbook(config.ARCHIVO_SAP)
        ws = wb[config.NOMBRE_HOJA]
        headers = {cell.value.strip(): cell.col_idx for cell in ws[1] if cell.value}
        count = 0
        
        for row in ws.iter_rows(min_row=2):
            val_cell = row[headers[config.COL_UBICACION]-1].value
            if not val_cell: continue
            
            val = str(val_cell).strip().upper()
            if val in resultados_finales:
                res = resultados_finales[val]
                ws.cell(row=row[0].row, column=headers[config.COL_STATUS_OUTPUT], value=res['status'])
                count += 1
                
        wb.save(config.ARCHIVO_SAP)
        print(f"[OK] SAP Actualizado: {count} filas.")
        return count
    except Exception as e:
        print(f"[ERROR] Error SAP final: {e}")
        return 0
